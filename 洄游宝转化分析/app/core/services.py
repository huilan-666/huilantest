import csv
import hashlib
import io
import json
import uuid
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException

from django.conf import settings
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from .models import Audit, Batch, ImportPreview, Record
from .policy import scoped

TZ = ZoneInfo('Asia/Shanghai')
FIELDS = {
    'name': ('商机名称', '客户名称', '企业名称'),
    'opportunity_id': ('商机ID', '商机编号', '销售机会ID'),
    'customer_id': ('客户ID', '客户编号'),
    'created': ('商机创建时间', '销售机会创建时间'),
    'department': ('所属部门', '部门'),
    'owner': ('跟进人', '负责人'),
    'source': ('线索来源', '来源'),
    'scenario': ('回流用途', '需求场景'),
    'stage': ('商机阶段', '阶段'),
    'quoted': ('报价时间', '报价日期'),
    'registered': ('注册时间', '注册日期'),
    'trade': ('首次交易时间', '客户历史首交时间'),
    'identity': ('新客确认', '客户类型', '新老客'),
    'identity_evidence': ('新客依据', '身份依据'),
    'unique_link': ('首交唯一关联', '唯一关联'),
    'last_followup': ('最近跟进时间',),
    'followup': ('最近跟进结果', '跟进记录'),
    'next_step': ('下一步安排',),
    'blocker': ('阻塞原因',),
}
FIELD_NAMES = {k: v[0] for k, v in FIELDS.items()}
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50000
MAX_EXPANDED = 200 * 1024 * 1024


class DomainError(Exception):
    def __init__(self, message, code='invalid', status=400):
        super().__init__(message)
        self.code, self.status = code, status


def timestamp(value):
    if value is None or str(value).strip() == '':
        return None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    else:
        try:
            parsed = datetime.fromisoformat(str(value).strip().replace('/', '-').replace('Z', '+00:00'))
        except ValueError as exc:
            raise DomainError('日期无法解析：' + str(value)[:100], 'date_parse') from exc
    return parsed.replace(tzinfo=TZ) if parsed.tzinfo is None else parsed.astimezone(TZ)


def iso(value):
    return timestamp(value).isoformat() if value not in [None, ''] else None


def date_value(value, label):
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError) as exc:
        raise DomainError(label + '必须为有效日期。') from exc


def config_values(data):
    start = date_value(data.get('start'), '创建起日')
    end = date_value(data.get('end'), '创建止日')
    cutoff = date_value(data.get('cutoff'), '结果截止日')
    if start > end or end > cutoff:
        raise DomainError('日期需满足：创建起日≤创建止日≤结果截止日。')
    exported = iso(data.get('exported_at'))
    if exported and timestamp(exported) < datetime.combine(cutoff, time.min, TZ):
        raise DomainError('导出时刻不能早于结果截止日。')
    return {'start': start.isoformat(), 'end': end.isoformat(), 'cutoff': cutoff.isoformat(), 'exported_at': exported, 'source_timezone': 'Asia/Shanghai'}


def stringify(value):
    return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat() if isinstance(value, date) else value


def read_file(content, filename, selected_sheet=None):
    if len(content) > MAX_BYTES:
        raise DomainError('文件超过20MiB上限，请缩小导出范围。', 'upload_limit', 413)
    ext = Path(filename).suffix.lower()
    try:
        if ext == '.csv':
            rows = list(csv.reader(io.StringIO(content.decode('utf-8-sig'))))
            sheets, sheet = ['CSV'], 'CSV'
        elif ext == '.xlsx':
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(x.file_size for x in z.infolist()) > MAX_EXPANDED:
                    raise DomainError('解压后文件超过200MiB上限。', 'expanded_limit', 413)
                for x in z.infolist():
                    parts = Path(x.filename.replace('\\', '/')).parts
                    if '..' in parts or x.filename.startswith(('/', '\\')) or ':' in x.filename:
                        raise DomainError('压缩包包含非法路径。', 'unsafe_file')
                    if 'vbaproject' in x.filename.lower():
                        raise DomainError('不接受包含宏的文件。', 'unsafe_file')
                    if x.filename.lower().endswith('.xml'):
                        # 禁止外部实体与DTD；解析引擎另由defusedxml防护。
                        xml = z.read(x)
                        if b'<!DOCTYPE' in xml.upper() or b'<!ENTITY' in xml.upper():
                            raise DomainError('文件含不安全XML定义。', 'unsafe_xml')
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            sheets = wb.sheetnames
            sheet = selected_sheet if selected_sheet in sheets else sheets[0]
            ws = wb[sheet]
            if ws.max_row and ws.max_row > MAX_ROWS + 1:
                raise DomainError('单表超过50000行上限。', 'row_limit', 413)
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS:
                    raise DomainError('单表超过50000行上限。', 'row_limit', 413)
                rows.append([stringify(v) for v in row])
            wb.close()
        else:
            raise DomainError('仅接受.xlsx和UTF-8 CSV，不接受宏文件。', 'file_type')
    except DomainError:
        raise
    except (ValueError, KeyError, OSError, UnicodeError, csv.Error, zipfile.BadZipFile, EOFError, ParseError, DefusedXmlException, InvalidFileException) as exc:
        raise DomainError('文件损坏或格式无效，请重新导出Excel或UTF-8 CSV。', 'file_parse') from exc
    if not rows or not any(str(v or '').strip() for v in rows[0]):
        raise DomainError('文件没有有效标题行。', 'empty_file')
    if len(rows) > MAX_ROWS + 1:
        raise DomainError('单表超过50000行上限。', 'row_limit', 413)
    headers = [str(v or '').strip() for v in rows[0]]
    nonempty = [h for h in headers if h]
    if len(nonempty) != len(set(nonempty)):
        raise DomainError('存在同名字段，请在CRM导出中使用唯一列名。', 'duplicate_header')
    raw = [{'row_number': i + 2, 'values': dict(zip(headers, row))} for i, row in enumerate(rows[1:]) if any(v not in [None, ''] for v in row)]
    if not raw:
        raise DomainError('文件只有标题，没有业务记录。', 'empty_file')
    return sheets, sheet, headers, raw


def guess_mapping(headers):
    return {k: next((v for v in aliases if v in headers), '') for k, aliases in FIELDS.items()}


def create_preview(user, upload, selected_sheet=None):
    content = upload.read(MAX_BYTES + 1)
    sheets, sheet, headers, raw = read_file(content, upload.name, selected_sheet)
    digest = hashlib.sha256(content).hexdigest()
    directory = Path(settings.MEDIA_ROOT)
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / (digest + Path(upload.name).suffix.lower())
    if not target.exists():
        temporary = directory / (uuid.uuid4().hex + '.tmp')
        try:
            temporary.write_bytes(content)
            temporary.replace(target)
        finally:
            if temporary.exists():
                temporary.unlink()
    created = []
    mapping = guess_mapping(headers)
    for row in raw:
        try:
            dt = timestamp(row['values'].get(mapping['created']))
            if dt:
                created.append(dt.date())
        except DomainError:
            continue
    start, end = (min(created), max(created)) if created else (date.today(), date.today())
    return ImportPreview.objects.create(owner=user, filename=Path(upload.name).name, file_hash=digest, file_path=str(target), sheets=sheets, sheet=sheet, headers=headers, raw_rows=raw, mapping=mapping, config={'start': str(start), 'end': str(end), 'cutoff': str(end), 'source_timezone': 'Asia/Shanghai'})


def normalized(preview, mapping, config, excluded=None, identity_overrides=None):
    excluded, identity_overrides = excluded or {}, identity_overrides or {}
    if not mapping.get('created') or not mapping.get('department'):
        raise DomainError('请映射商机创建时间和所属部门，缺少这些字段无法发布。', 'missing_field')
    if any(v and v not in preview.headers for v in mapping.values()):
        raise DomainError('映射包含不存在的字段。', 'missing_field')
    result = []
    for source in preview.raw_rows:
        number, raw = source['row_number'], source['values']
        data = {k: raw.get(mapping.get(k), '') for k in FIELDS}
        issues = []
        for key in ['created', 'quoted', 'registered', 'trade', 'last_followup']:
            try:
                data[key] = iso(data[key])
            except DomainError:
                data[key] = None
                issues.append(FIELD_NAMES[key] + '无法解析')
        data['department'] = str(data['department'] or '').strip()
        for k in ['opportunity_id', 'customer_id', 'name', 'owner', 'source', 'scenario', 'stage', 'identity_evidence', 'followup', 'next_step', 'blocker']:
            data[k] = str(data[k] or '').strip()
        if not data['name']:
            data['name'] = '商机第' + str(number) + '行'
        if not data['created']:
            issues.append('缺有效商机创建时间')
        if not data['department']:
            issues.append('缺所属部门')
        created = timestamp(data['created'])
        for k in ['quoted', 'registered', 'trade']:
            event = timestamp(data[k])
            if event and created and event < created:
                issues.append(FIELD_NAMES[k] + '早于商机创建')
            if event and created and (event-created).total_seconds() > 100*86400:
                issues.append(FIELD_NAMES[k] + '超过100天，保留分析')
        if data['trade'] and data['registered'] and timestamp(data['trade']) < timestamp(data['registered']):
            issues.append('客户首交早于注册')
        identity_input = str(data.get('identity') or '').strip().lower()
        identity = 'new' if identity_input in ['新客', '是', 'yes', 'true', 'new', '1'] else 'existing' if identity_input in ['老客', '否', 'existing', 'no', 'false', '0'] else 'unknown'
        if data['trade'] and created and timestamp(data['trade']) < created:
            identity = 'existing'
        unique = str(data.get('unique_link') or '').strip().lower() in ['是', 'yes', 'true', '1', '已确认']
        override = identity_overrides.get(str(number), {})
        if override:
            data['opportunity_id'] = str(override.get('opportunity_id', data['opportunity_id'])).strip()
            data['customer_id'] = str(override.get('customer_id', data['customer_id'])).strip()
            identity, unique = override['identity'], override['unique_link']
            data['identity_evidence'] = override['evidence']
            if data['trade'] and created and timestamp(data['trade']) < created and identity == 'new':
                raise DomainError('第' + str(number) + '行已有早于商机的客户首交，不能确认为新客。')
        if identity == 'new' and not data['identity_evidence']:
            identity = 'unknown'
            issues.append('新客缺确认依据')
        exclusion = excluded.get(str(number), '')
        if not exclusion and (not created or not data['department']):
            issues.append('发布前须修正或明确剔除')
        result.append({'row_number': number, 'raw': raw, 'data': data, 'department': data['department'], 'opportunity_id': data['opportunity_id'], 'customer_id': data['customer_id'], 'identity': identity, 'identity_evidence': data['identity_evidence'], 'unique_link': unique, 'excluded': bool(exclusion), 'exclusion_reason': exclusion, 'issues': issues})
    # 同一客户只有一个有依据的新客首交商机；冲突不能静默取第一行。
    ids = Counter(r['opportunity_id'] for r in result if r['opportunity_id'] and not r['excluded'])
    customer_links = Counter(r['customer_id'] for r in result if r['customer_id'] and r['unique_link'] and r['identity'] == 'new' and not r['excluded'])
    for r in result:
        if r['opportunity_id'] and ids[r['opportunity_id']] > 1:
            r['unique_link'] = False
            r['issues'].append('商机ID重复，关联待核实')
        if r['customer_id'] and customer_links[r['customer_id']] > 1:
            r['unique_link'] = False
            r['issues'].append('同客户多个首交关联，需指定唯一商机')
    return result


def counts(rows, config):
    start, end = date_value(config['start'], '起日'), date_value(config['end'], '止日')
    cutoff = datetime.combine(date_value(config['cutoff'], '截止') + timedelta(days=1), time.min, TZ)
    in_range = [r for r in rows if r['data']['created'] and start <= timestamp(r['data']['created']).date() <= end]
    valid = [r for r in in_range if not r['excluded']]
    return {'raw': len(rows), 'out_of_range': len(rows)-len(in_range), 'excluded': sum(r['excluded'] for r in in_range), 'valid': len(valid), 'unknown': sum(r['identity']=='unknown' for r in valid), 'missing_ids': sum(not r['opportunity_id'] or not r['customer_id'] for r in valid), 'future_events': sum(timestamp(r['data'][k]) >= cutoff for r in rows for k in ['quoted','registered','trade'] if r['data'][k]), 'issues': sum(bool(r['issues']) for r in rows)}


def publish(preview, user, mapping, config, exclusions, overrides=None, previous=None):
    if any(not str(v).strip() for v in exclusions.values()):
        raise DomainError('每条明确剔除都需要理由。')
    rows = normalized(preview, mapping, config, exclusions, overrides)
    if any('发布前须修正或明确剔除' in r['issues'] for r in rows):
        raise DomainError('存在缺有效创建时间或所属部门的记录，请修正或明确剔除。', 'required_row')
    key = json.dumps([preview.file_hash, preview.sheet, mapping, config, exclusions, overrides or {}], sort_keys=True, ensure_ascii=False)
    fingerprint = hashlib.sha256(key.encode()).hexdigest()
    with transaction.atomic():
        existing = Batch.objects.filter(fingerprint=fingerprint).first()
        if existing:
            return existing, False
        prior = previous or Batch.objects.filter(file_hash=preview.file_hash).order_by('-version').first()
        batch = Batch.objects.create(owner=user, name=preview.filename, fingerprint=fingerprint, file_hash=preview.file_hash, source_path=preview.file_path, sheet=preview.sheet, config=config, mapping=mapping, counts=counts(rows, config), version=(prior.version+1 if prior else 1), previous=prior)
        Record.objects.bulk_create([Record(batch=batch, **r) for r in rows])
        Audit.objects.create(author=user, entity='batch', entity_id=batch.pk, event='publish', version=batch.version, detail={'fingerprint': fingerprint, 'counts': batch.counts})
    return batch, True


def visible_records(batch, user):
    return scoped(batch.records.all(), user)


def record_ready(r):
    return r.identity == 'new' and bool(r.opportunity_id and r.customer_id and r.unique_link and r.identity_evidence)


def stage_advice(record):
    created = timestamp(record.data.get('created'))
    if not created:
        return {'candidate': '创建时间缺失，无法判断推进阶段与时间顺序。', 'task': '核对源文件的商机创建时间和字段映射，修正后生成新快照。', 'evidence': '原始创建时间及字段映射依据'}
    cutoff = datetime.combine(date_value(record.batch.config['cutoff'], '截止日') + timedelta(days=1), time.min, TZ)
    def occurred(key):
        event = timestamp(record.data.get(key))
        return bool(event and created <= event < cutoff)
    if record.identity == 'existing':
        return {'candidate':'已有早于当前商机的客户历史首交，不能用于解释本次新客开户。','task':'先核实新老客与商机关联，老客本次交易不在本期统计范围。','evidence':'客户完整交易历史及本商机的关联依据'}
    if occurred('trade'):
        return {'candidate':'截止日前已记录客户首交，可作为同龄商机的推进对照；不能据此推断某种动作有效。','task':'核对首交关联和推进证据，整理可供业务复核的经验。','evidence':'首交唯一关联及真实沟通过程'}
    if occurred('registered'):
        return {'candidate':'已注册未首交：合同签约、资金需求、方案顾虑或关键决策承诺等仍待核实，现有日期不能确定原因。','task':'确认当前首测计划、合同及决策反馈，核实具体阻塞，约定下一步并记录客户反馈。','evidence':'日期、金额币种、用途、KP确认及注册后真实客户反馈'}
    if occurred('quoted'):
        return {'candidate':'已报价未注册：报价反馈、决策安排、客户计划时点或资料准备情况待核实。','task':'核实报价接受度、关键决策人安排及注册资料状态，明确下一次推进动作。','evidence':'报价反馈、决策链与资料准备记录'}
    return {'candidate':'未记录报价：可能仍在挖需/方案沟通、客户暂缓，或报价未及时更新；不能把空日期当未执行。','task':'核实最近沟通、具体产品需求、报价是否已发送及下一次沟通计划，再决定是否调整跟进或辅导。','evidence':'最近跟进结果、需求确认、实际报价与下一步安排'}


def report(batch, user, window=None, group='stage', start=None, end=None, cutoff_date=None):
    config = batch.config
    start = date_value(start or config['start'], '起日')
    end = date_value(end or config['end'], '止日')
    cutoff_day = date_value(cutoff_date or config['cutoff'], '截止日')
    if start>end or end>cutoff_day or cutoff_day>date_value(config['cutoff'],'快照截止'):
        raise DomainError('分析范围无效，结果截止不可晚于快照截止日。')
    if window not in [None, '']:
        try:
            window=int(window)
        except (ValueError,TypeError) as exc:
            raise DomainError('观察窗口必须为正整数天。') from exc
        if not 1<=window<=3650:
            raise DomainError('观察窗口须为1至3650天。')
    else:
        window=None
    cutoff = datetime.combine(cutoff_day+timedelta(days=1),time.min,TZ)
    records=[r for r in visible_records(batch,user).order_by('row_number') if not r.excluded and r.data['created'] and start<=timestamp(r.data['created']).date()<=end]
    known_old=[r for r in records if r.identity=='existing']
    potential=[r for r in records if r.identity!='existing']
    formal=bool(potential) and all(record_ready(r) for r in potential)
    records=potential if formal else records
    mature=[r for r in records if not window or timestamp(r.data['created'])+timedelta(days=window)<cutoff]
    def has_event(r,key):
        event=timestamp(r.data.get(key))
        created=timestamp(r.data['created'])
        return bool(event and created<=event<cutoff and (not window or event-created<=timedelta(days=window)))
    n=len(mature)
    mature_ids={r.pk for r in mature}
    events={k:sum(has_event(r,k) for r in mature) for k in ['quoted','registered','trade']}
    stages=Counter()
    group_rows=defaultdict(list)
    details=[]
    for r in records:
        # 进展分布使用快照截止时点，不随固定窗口定义混淆。
        stage='trade' if r.data.get('trade') and timestamp(r.data['created'])<=timestamp(r.data['trade'])<cutoff else 'registered' if r.data.get('registered') and timestamp(r.data['registered'])<cutoff else 'quoted' if r.data.get('quoted') and timestamp(r.data['quoted'])<cutoff else 'unquoted'
        stages[stage]+=1
        age=(cutoff-timestamp(r.data['created'])).total_seconds()/86400
        key={'stage':{'trade':'已有首交','registered':'注册未首交','quoted':'报价未注册','unquoted':'未记录报价'}[stage], 'source':r.data.get('source') or '未记录来源','department':r.department,'scenario':r.data.get('scenario') or '未记录场景'}.get(group)
        if key is None:
            raise DomainError('不支持的分组维度。')
        group_rows[key].append(r)
        details.append({'id':r.pk,'row':r.row_number,'name':r.data['name'],'department':r.department,'owner':r.data.get('owner'),'source':r.data.get('source'),'scenario':r.data.get('scenario'),'stage':stage,'lost':'丢单' in r.data.get('stage',''),'age':round(age,1),'mature':r.pk in mature_ids,'issues':r.issues,'identity':r.identity,'local_only':not(r.opportunity_id and r.customer_id),'ready':record_ready(r)})
    groups=[]
    for key,rs in group_rows.items():
        eligible=[r for r in rs if r.pk in mature_ids]
        success=sum(has_event(r,'trade') for r in eligible)
        groups.append({'label':key,'count':len(rs),'denominator':len(eligible),'numerator':success,'rate':round(success/len(eligible)*100,2) if eligible else None,'share':round(len(rs)/len(records)*100,1) if records else 0,'mean_age':round(sum((cutoff-timestamp(r.data['created'])).total_seconds()/86400 for r in rs)/len(rs),1)})
    return {'batch_id':batch.pk,'version':batch.version,'config':{'start':str(start),'end':str(end),'cutoff':str(cutoff_day),'window':window},'mode':'formal' if formal else 'sample','total':len(records),'denominator':n,'numerator':events['trade'],'rate':round(events['trade']/n*100,2) if n else None,'quoted':events['quoted'],'registered':events['registered'],'immature':len(records)-n,'unknown':sum(not record_ready(r) for r in potential),'old':len(known_old),'stages':dict(stages),'groups':groups,'records':details,'generated_at':timezone.now().isoformat(),'limitation':'当前属性分组，不用于历史责任判断；观察差异不证明原因。'}
